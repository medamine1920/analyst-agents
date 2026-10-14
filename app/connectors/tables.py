"""Turning what a service returns into tables (CSV, Excel, JSON, MCP results), and tables into files.

A table is {"name": str, "rows": [dict, ...]}. Values are flattened to text or numbers, so any table
can become a CSV file the data-source builder understands.
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

MAX_ROWS = 200_000
MAX_COLUMNS = 200


def flatten(value: Any) -> Any:
    """Nested values (lists, objects) become compact text; numbers and text stay as they are."""
    if value is None or isinstance(value, (int, float, str, bool)):
        return value
    if isinstance(value, list) and all(isinstance(v, (str, int, float)) for v in value):
        return ", ".join(str(v) for v in value)
    return json.dumps(value, default=str, ensure_ascii=False)


def table(name: str, rows: list[dict]) -> dict:
    rows = [{str(k)[:100]: flatten(v) for k, v in row.items()} for row in rows[:MAX_ROWS] if isinstance(row, dict)]
    columns = list(dict.fromkeys(k for row in rows for k in row))[:MAX_COLUMNS]
    return {"name": name, "rows": [{c: row.get(c) for c in columns} for row in rows], "columns": columns}


def from_csv(name: str, text: str) -> dict:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return table(name, list(csv.DictReader(io.StringIO(text), dialect=dialect)))


def from_excel(name: str, content: bytes) -> list[dict]:
    """Every sheet with a header row becomes a table."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("This file could not be read as an Excel workbook (.xlsx).") from exc
    tables = []
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not any(header):
            continue
        names = [str(h) if h is not None else f"column_{i + 1}" for i, h in enumerate(header)]
        body = [dict(zip(names, row, strict=False)) for row in rows if any(v is not None for v in row)]
        tables.append(table(f"{name} - {sheet.title}" if len(book.worksheets) > 1 else name, body))
    if not tables:
        raise ValueError("The workbook has no sheet with a header row.")
    return tables


def records_in(payload: Any) -> list[tuple[str, list[dict]]]:
    """Find lists of records in JSON: the payload itself, or lists under its keys (one level down)."""
    if isinstance(payload, list) and payload and all(isinstance(r, dict) for r in payload):
        return [("data", payload)]
    found = []
    if isinstance(payload, dict):
        for key, value in payload.items():
            if isinstance(value, list) and value and all(isinstance(r, dict) for r in value):
                found.append((key, value))
        if not found and "columns" in payload and "rows" in payload:  # {"columns": [...], "rows": [[...]]}
            found.append(("data", [dict(zip(payload["columns"], r, strict=False)) for r in payload["rows"]]))
    return found


def from_json(name: str, payload: Any) -> list[dict]:
    found = records_in(payload)
    if not found:
        raise ValueError("No list of records was found in this JSON.")
    return [table(name if len(found) == 1 else f"{name} - {key}", rows) for key, rows in found]


def from_bytes(name: str, content: bytes, content_type: str = "", filename: str = "") -> list[dict]:
    """Guess the format from the name, the content type, then the content itself."""
    lowered = (filename or "").lower()
    if lowered.endswith(".xlsx") or "spreadsheetml" in content_type or content[:2] == b"PK":
        return from_excel(name, content)
    text = content.decode("utf-8-sig", errors="replace")
    if lowered.endswith(".json") or "json" in content_type or text.lstrip()[:1] in ("[", "{"):
        try:
            return from_json(name, json.loads(text))
        except json.JSONDecodeError:
            pass
    if lowered.endswith((".csv", ".tsv", ".txt")) or "csv" in content_type or "text/plain" in content_type or "," in text[:2000]:
        return [from_csv(name, text)]
    raise ValueError("This is not a CSV, Excel or JSON file.")


def to_csv(rows: list[dict]) -> bytes:
    columns = list(dict.fromkeys(k for row in rows for k in row))
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({c: _safe_cell(row.get(c)) for c in columns})
    return out.getvalue().encode("utf-8")


def _safe_cell(value: Any) -> Any:
    """A text cell starting with = + - @ would run as a formula when the file is opened in a spreadsheet."""
    if isinstance(value, str) and re.match(r"^[=+\-@]", value):
        return "'" + value
    return value


def to_excel(tables: list[dict]) -> bytes:
    """One sheet per table, header in bold, safe against formula injection."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for item in tables or [{"name": "Results", "rows": []}]:
        title = re.sub(r"[\[\]:*?/\\]", " ", item["name"])[:31] or "Sheet"
        while title in used:
            title = title[:28] + f"_{len(used)}"
        used.add(title)
        sheet = book.create_sheet(title)
        columns = list(dict.fromkeys(k for row in item["rows"] for k in row))
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in item["rows"]:
            sheet.append([row.get(c) for c in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                    cell.data_type = "s"  # stored as text: a spreadsheet will not run it as a formula
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
