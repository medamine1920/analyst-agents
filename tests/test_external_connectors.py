"""Connectors to outside tools, with the network faked: safety checks, formats, each kind, and privacy."""

import asyncio
import io
import json
import time

import httpx
import pytest
from fastapi.testclient import TestClient
from mcp.server import MCPServer
from mcp.server.transport_security import TransportSecuritySettings

from app.api.main import create_app
from app.api.security import SecurityStore
from app.connectors import http, kinds
from app.connectors.http import ConnectorError
from app.connectors.store import ConnectionStore
from app.connectors.tables import from_bytes, from_excel, to_excel
from app.mcp_server.server import server
from tests.test_api import plan_fn, write_fn

PASSWORD = "correct horse battery"
SALES_CSV = b"day,region,amount\n2025-01-02,North,10\n2025-01-03,South,4.5\n2025-02-01,North,7\n"


def run(coro):
    return asyncio.run(coro)


@pytest.fixture(autouse=True)
def fake_network(monkeypatch):
    """Names ending in .internal stand for private addresses; everything else is 'public'."""
    def check_host(host):
        if host.endswith(".internal") or host in ("127.0.0.1", "localhost"):
            raise ValueError(f"'{host}' is inside a private network, which this console does not connect to.")
    monkeypatch.setattr(http, "check_host", check_host)
    monkeypatch.setattr(http, "RETRIES", 2)
    monkeypatch.setattr(asyncio, "sleep", _no_wait)
    yield
    http.TRANSPORT = None


async def _no_wait(*_):
    return None


def serve(handler):
    http.TRANSPORT = httpx.MockTransport(handler)


# ---- the safe HTTP layer -------------------------------------------------------------------------

def test_private_addresses_are_refused_even_behind_a_redirect():
    with pytest.raises(ConnectorError, match="private network"):
        run(http.request("GET", "https://db.internal/data.csv"))
    serve(lambda request: httpx.Response(302, headers={"location": "https://admin.internal/secrets"}))
    with pytest.raises(ConnectorError, match="private network"):
        run(http.request("GET", "https://files.example.com/report.csv"))


def test_busy_services_are_retried_and_refusals_explained():
    calls = []

    def handler(request):
        calls.append(1)
        return httpx.Response(503) if len(calls) < 3 else httpx.Response(200, text="ok")
    serve(handler)
    assert run(http.request("GET", "https://api.example.com/x")).text == "ok" and len(calls) == 3
    serve(lambda request: httpx.Response(401))
    with pytest.raises(ConnectorError, match="refused the credentials"):
        run(http.request("GET", "https://api.example.com/x"))
    serve(lambda request: httpx.Response(200, content=b"x" * 2000))
    with pytest.raises(ConnectorError, match="larger than"):
        run(http.request("GET", "https://api.example.com/big", max_bytes=1000))
    with pytest.raises(ConnectorError, match="https://"):
        run(http.request("GET", "file:///etc/passwd"))


# ---- formats -------------------------------------------------------------------------------------

def test_csv_json_and_excel_all_become_tables():
    assert from_bytes("Sales", SALES_CSV, "text/csv")[0]["rows"][1] == {"day": "2025-01-03", "region": "South", "amount": "4.5"}
    semicolons = from_bytes("EU", b"day;amount\n2025-01-02;10\n", "", "eu.csv")[0]
    assert semicolons["columns"] == ["day", "amount"]
    nested = from_bytes("API", json.dumps({"results": [{"id": 1, "tags": ["a", "b"], "owner": {"name": "Ana"}}]}).encode())
    assert nested[0]["rows"][0] == {"id": 1, "tags": "a, b", "owner": '{"name": "Ana"}'}
    workbook = to_excel([{"name": "Sales", "rows": [{"region": "North", "amount": 10}, {"region": "=HYPERLINK(1)", "amount": 2}]}])
    back = from_excel("Book", workbook)[0]
    assert back["rows"][0] == {"region": "North", "amount": 10} and back["rows"][1]["region"] == "=HYPERLINK(1)"
    from openpyxl import load_workbook
    cell = load_workbook(io.BytesIO(workbook))["Sales"]["A3"]
    assert cell.data_type == "s"  # stored as text, never run as a formula
    with pytest.raises(ValueError):
        from_bytes("Bad", b"\x00\x01binary", "application/octet-stream", "x.bin")


# ---- each kind ---------------------------------------------------------------------------------

def test_a_shared_google_sheet_link_is_read_as_a_workbook():
    seen = []

    def handler(request):
        seen.append(str(request.url))
        return httpx.Response(200, content=to_excel([{"name": "Q1", "rows": [{"a": 1}]}, {"name": "Q2", "rows": [{"a": 2}]}]),
                              headers={"content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    serve(handler)
    tables = run(kinds.Link().fetch({}, {}, None, {"url": "https://docs.google.com/spreadsheets/d/abc123/edit#gid=0", "name": "Plan"}))
    assert seen[0] == "https://docs.google.com/spreadsheets/d/abc123/export?format=xlsx"
    assert [t["name"] for t in tables] == ["Plan - Q1", "Plan - Q2"]
    serve(lambda request: httpx.Response(200, text="<html>sign in</html>", headers={"content-type": "text/html"}))
    with pytest.raises(ConnectorError, match="not shared"):
        run(kinds.Link().fetch({}, {}, None, {"url": "https://docs.google.com/spreadsheets/d/abc123/edit"}))


def test_slack_and_teams_only_accept_their_own_webhook_addresses_and_get_the_answer():
    posted = []
    serve(lambda request: posted.append(json.loads(request.content)) or httpx.Response(200, text="ok"))
    payload = {"question": "Which store?", "answer": "Philadelphia had more.", "tables": [{"name": "By store", "rows": [{"store": "Philadelphia", "orders": 39931}]}]}
    with pytest.raises(ConnectorError, match="hooks.slack.com"):
        run(kinds.Slack().send({}, {"webhook_url": "https://evil.example.com/hook"}, payload))
    assert run(kinds.Slack().send({}, {"webhook_url": "https://hooks.slack.com/services/T/B/x"}, payload)) == "Sent to Slack."
    assert "Philadelphia | 39931" in json.dumps(posted[0])
    with pytest.raises(ConnectorError):
        run(kinds.Teams().send({}, {"webhook_url": "https://example.com/hook"}, payload))
    assert run(kinds.Teams().send({}, {"webhook_url": "https://prod-1.westeurope.logic.azure.com/workflows/x"}, payload)) == "Sent to Teams."


def test_airtable_notion_and_hubspot_follow_every_page():
    def airtable(request):
        if "/meta/bases/" in request.url.path:
            return httpx.Response(200, json={"tables": [{"id": "tbl1", "name": "Deals"}]})
        if request.url.params.get("offset") == "p2":
            return httpx.Response(200, json={"records": [{"id": "r2", "fields": {"amount": 5}}]})
        return httpx.Response(200, json={"records": [{"id": "r1", "fields": {"amount": 3}}], "offset": "p2"})
    serve(airtable)
    assert run(kinds.Airtable().items({"base_id": "app1"}, {"token": "t"})) == [{"id": "tbl1", "label": "Deals"}]
    assert [r["amount"] for r in run(kinds.Airtable().fetch({"base_id": "app1"}, {"token": "t"}, "tbl1", {}))[0]["rows"]] == [3, 5]

    def notion(request):
        body = json.loads(request.content or b"{}")
        page = {"properties": {"Name": {"type": "title", "title": [{"plain_text": "Acme"}]},
                               "Stage": {"type": "select", "select": {"name": "Won"}}, "Value": {"type": "number", "number": 9}}}
        more = not body.get("start_cursor")
        return httpx.Response(200, json={"results": [page], "has_more": more, "next_cursor": "c2" if more else None})
    serve(notion)
    rows = run(kinds.Notion().fetch({}, {"token": "t"}, "db1", {}))[0]["rows"]
    assert rows == [{"Name": "Acme", "Stage": "Won", "Value": 9}] * 2

    def hubspot(request):
        if "/properties/" in request.url.path:
            return httpx.Response(200, json={"results": [{"name": "dealname"}, {"name": "secret_calc", "calculated": True}]})
        after = request.url.params.get("after")
        return httpx.Response(200, json={"results": [{"id": after or "1", "properties": {"dealname": "Big"}}],
                                         **({} if after else {"paging": {"next": {"after": "2"}}})})
    serve(hubspot)
    rows = run(kinds.HubSpot().fetch({}, {"token": "t"}, "deals", {}))[0]["rows"]
    assert [r["id"] for r in rows] == ["1", "2"] and "secret_calc" not in rows[0]


def test_any_mcp_server_lists_its_tools_and_a_tool_with_records_is_imported():
    remote = MCPServer(name="crm")

    @remote.tool(description="Monthly sales")
    def sales() -> list[dict]:
        return [{"month": "2025-01", "amount": 10}, {"month": "2025-02", "amount": 12}]

    app = remote.streamable_http_app(streamable_http_path="/mcp", stateless_http=True, json_response=True,
                                     transport_security=TransportSecuritySettings(enable_dns_rebinding_protection=False))

    async def scenario():
        http.TRANSPORT = httpx.ASGITransport(app=app)
        async with app.router.lifespan_context(app):
            settings = {"url": "https://crm.example.com/mcp"}
            tools = await kinds.Mcp().items(settings, {})
            tables = await kinds.Mcp().fetch(settings, {}, "sales", {})
        return tools, tables

    tools, tables = run(scenario())
    assert [t["id"] for t in tools] == ["sales"] and tables[0]["rows"][1] == {"month": "2025-02", "amount": 12}
    with pytest.raises(ConnectorError, match="private network"):
        run(kinds.Mcp().items({"url": "https://tools.internal/mcp"}, {}))


def test_credentials_are_encrypted_and_shown_only_as_a_hint(tmp_path):
    store = ConnectionStore(tmp_path / "console.sqlite")
    saved = store.add("ana", "airtable", "CRM", {"base_id": "app1"}, {"token": "patSECRET1234567890"})
    assert "patSECRET1234567890" not in json.dumps(saved) and saved["hint"].startswith("token: patS…")
    assert b"patSECRET1234567890" not in (tmp_path / "console.sqlite").read_bytes()
    assert store.secret("ana", saved["id"]) == {"token": "patSECRET1234567890"}
    with pytest.raises(KeyError):
        store.secret("bob", saved["id"])  # private to its owner
    state, _ = store.new_state("ana", "google", "Drive")
    assert store.take_state(state)["owner"] == "ana" and store.take_state(state) is None  # one use only


# ---- through the API ------------------------------------------------------------------------------

@pytest.fixture
def api(tmp_path):
    accounts = SecurityStore(tmp_path / "console.sqlite")
    accounts.add_user("ana", PASSWORD)
    accounts.add_user("bob", PASSWORD)
    app = create_app(mcp_target=server, plan_fn=plan_fn, write_fn=write_fn, traces_dir=tmp_path / "traces",
                     checkpoint_db=tmp_path / "c.sqlite", results_dir=tmp_path, auth=True,
                     security_db=tmp_path / "console.sqlite", sources_dir=tmp_path / "sources")
    with TestClient(app) as client:
        yield client


def sign_in(client, name):
    client.cookies.clear()
    client.post("/api/auth/login", json={"name": name, "password": PASSWORD})


@pytest.mark.integration
def test_a_link_import_becomes_a_data_source_and_results_export_to_excel(api):
    serve(lambda request: httpx.Response(200, content=SALES_CSV, headers={"content-type": "text/csv"}))
    sign_in(api, "ana")
    created = api.post("/api/connectors/link/import", json={"url": "https://files.example.com/sales.csv", "name": "Sales"}).json()
    assert created["tables_found"] == [{"name": "Sales", "rows": 3}]
    for _ in range(240):
        source = api.get(f"/api/sources/{created['id']}").json()
        if source["status"] != "building":
            break
        time.sleep(1)
    assert source["status"] == "ready", source.get("error")
    excel = api.post("/api/export/xlsx", json={"name": "Q1", "tables": [{"name": "Sales", "rows": [{"a": 1}]}]})
    assert excel.status_code == 200 and excel.content[:2] == b"PK"
    catalog = api.get("/api/connectors/catalog").json()
    assert {"link", "mcp", "google", "slack", "teams", "airtable", "notion", "hubspot"} <= {k["kind"] for k in catalog["kinds"]}
    assert "linkedin" in catalog["not_possible"]
    assert api.post("/api/connectors", json={"kind": "linkedin", "values": {}}).status_code == 404


@pytest.mark.integration
def test_saved_connections_never_return_their_secret_and_runs_stay_private(api):
    serve(lambda request: httpx.Response(200, json={"tables": [{"id": "tbl1", "name": "Deals"}]}))
    sign_in(api, "ana")
    saved = api.post("/api/connectors", json={"kind": "airtable", "name": "CRM",
                                              "values": {"base_id": "app1", "token": "patSECRET1234567890"}}).json()
    listed = api.get("/api/connectors").text
    assert "patSECRET1234567890" not in listed and "patSECRET1234567890" not in json.dumps(saved)
    assert api.get(f"/api/connectors/{saved['id']}/items").json() == [{"id": "tbl1", "label": "Deals"}]
    refused = api.post("/api/connectors", json={"kind": "slack", "values": {"webhook_url": "https://evil.example.com/x"}})
    assert refused.status_code == 422

    events = api.post("/api/ask", json={"question": "How many orders were placed?"}).text
    trace_id = json.loads(next(line for line in events.splitlines() if '"saved"' in line)[6:])["id"]
    assert api.get(f"/api/traces/{trace_id}").status_code == 200
    sign_in(api, "bob")
    assert api.get(f"/api/traces/{trace_id}").status_code == 404  # someone else's run
    assert api.get(f"/api/connectors/{saved['id']}/items").status_code == 404
    assert trace_id not in {t["id"] for t in api.get("/api/traces").json()}
